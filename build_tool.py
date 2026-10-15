# -*- coding: utf-8 -*-
"""生成「组件 BUG 统计工具」离线 HTML + 预导出 Excel。"""
import json, os, sys
from collections import Counter, defaultdict
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

HERE = os.path.dirname(os.path.abspath(__file__))
MAPPING_FILE = sys.argv[1] if len(sys.argv) > 1 else "/Users/weizhuangziji/Downloads/组件对应模块明细.xlsx"
BUG_FILE = sys.argv[2] if len(sys.argv) > 2 else None
XLSX_LIB = "/tmp/xlsx.full.min.js"
ECHARTS_LIB = "/tmp/echarts.min.js"
SEV_LEVELS = {"严重", "紧急"}


def load_mapping(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.worksheets[0]
    order, mapping = [], {}
    for r in range(2, ws.max_row + 1):
        comp = ws.cell(r, 1).value
        if comp is None:
            continue
        comp = str(comp).strip()
        mod = ws.cell(r, 2).value
        order.append(comp)
        mapping[comp] = str(mod).strip() if mod is not None else ""
    wb.close()
    return order, mapping


def load_bug_records(path):
    """读取 BUG 表：D 列“优先级” + M 列“组件”。"""
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.worksheets[0]
    headers = [str(ws.cell(1, c).value).strip() if ws.cell(1, c).value is not None else "" for c in range(1, ws.max_column + 1)]
    cidx = next((i for i, h in enumerate(headers) if h == "组件"), None)
    pidx = next((i for i, h in enumerate(headers) if h == "优先级"), None)
    if cidx is None:
        raise ValueError("未找到“组件”列")
    records = []
    for r in range(2, ws.max_row + 1):
        comp = ws.cell(r, cidx + 1).value
        if comp is None or not str(comp).strip():
            continue
        pri = ws.cell(r, pidx + 1).value if pidx is not None else None
        records.append({"comp": str(comp).strip(), "priority": str(pri).strip() if pri is not None else ""})
    wb.close()
    return records


def split_tokens(s):
    return [p.strip() for p in str(s).replace("，", ",").replace("；", ",").replace(";", ",").split(",") if p.strip()]


def compute(order, mapping, bug_records):
    counts = Counter()
    sev_comp = Counter()
    sev_mod = Counter()
    severe_urgent_total = 0
    for rec in bug_records:
        toks = split_tokens(rec["comp"])
        for tok in toks:
            counts[tok] += 1
        if rec["priority"] in SEV_LEVELS:
            severe_urgent_total += 1
            for tok in toks:
                sev_comp[tok] += 1
                sev_mod[mapping.get(tok, "未分类")] += 1

    components = [{"component": c, "module": mapping.get(c, "未分类"), "bugCount": counts.get(c, 0), "severeUrgent": sev_comp.get(c, 0)} for c in order]
    for tok, cnt in counts.items():
        if tok not in mapping:
            components.append({"component": tok, "module": "未分类", "bugCount": cnt, "severeUrgent": sev_comp.get(tok, 0)})
    components.sort(key=lambda x: (-x["bugCount"], x["component"]))

    mod_bug, mod_comp = defaultdict(int), defaultdict(int)
    for c in components:
        mod_bug[c["module"]] += c["bugCount"]
        if c["bugCount"] > 0:
            mod_comp[c["module"]] += 1
    modules = [{"module": m, "bugCount": mod_bug[m], "componentCount": mod_comp[m], "severeUrgent": sev_mod.get(m, 0)} for m in mod_bug]
    modules.sort(key=lambda x: (-x["bugCount"], x["module"]))
    return modules, components, severe_urgent_total


def write_excel(path, modules, components, bug_total):
    wb = openpyxl.Workbook()
    head_fill = PatternFill("solid", fgColor="2563EB")
    head_font = Font(color="FFFFFF", bold=True)
    ws1 = wb.active
    ws1.title = "模块汇总"
    ws1.append(["#", "模块", "有BUG组件数", "BUG数量", "严重+紧急", "占比"])
    for i, m in enumerate(modules, 1):
        ws1.append([i, m["module"], m["componentCount"], m["bugCount"], m["severeUrgent"], f"{m['bugCount']/bug_total*100:.2f}%" if bug_total else "0.00%"])
    ws2 = wb.create_sheet("组件明细")
    ws2.append(["#", "组件", "模块", "BUG数量", "占比"])
    for i, c in enumerate(components, 1):
        ws2.append([i, c["component"], c["module"], c["bugCount"], f"{c['bugCount']/bug_total*100:.2f}%" if bug_total else "0.00%"])
    for ws in (ws1, ws2):
        for cell in ws[1]:
            cell.fill = head_fill
            cell.font = head_font
            cell.alignment = Alignment(horizontal="center")
        ws.freeze_panes = "A2"
    for col, width in zip("ABCDEF", (6, 20, 14, 12, 12, 10)):
        ws1.column_dimensions[col].width = width
    for col, width in zip("ABCDE", (6, 22, 14, 12, 10)):
        ws2.column_dimensions[col].width = width
    wb.save(path)


def main():
    order, mapping = load_mapping(MAPPING_FILE)
    bug_records = load_bug_records(BUG_FILE) if BUG_FILE else []
    modules, components, severe_urgent_total = compute(order, mapping, bug_records)
    bug_total = len(bug_records)

    with open(os.path.join(HERE, "template.html"), encoding="utf-8") as f:
        html = f.read()
    with open(XLSX_LIB, encoding="utf-8") as f:
        xlsx_lib = f.read().replace("</script", "<\\/script")
    with open(ECHARTS_LIB, encoding="utf-8") as f:
        echarts_lib = f.read().replace("</script", "<\\/script")

    html = (html
            .replace("__XLSX_LIB__", xlsx_lib)
            .replace("__ECHARTS_LIB__", echarts_lib)
            .replace("__MAPPING_ORDER__", json.dumps(order, ensure_ascii=False))
            .replace("__MAPPING_MAP__", json.dumps(mapping, ensure_ascii=False))
            .replace("__BUG_RECORDS__", json.dumps(bug_records, ensure_ascii=False)))

    html_path = os.path.join(HERE, "组件BUG统计工具.html")
    with open(html_path, "w", encoding="utf-8") as f:
        f.write(html)
    xlsx_path = os.path.join(HERE, "组件BUG统计_结果.xlsx")
    if BUG_FILE:
        write_excel(xlsx_path, modules, components, bug_total)

    print("BUG 记录总数:", bug_total)
    print("严重+紧急:", severe_urgent_total)
    print("组件计数合计:", sum(c["bugCount"] for c in components))
    print("涉及组件数:", sum(1 for c in components if c["bugCount"] > 0))
    print("涉及模块数:", sum(1 for m in modules if m["bugCount"] > 0))
    print("生成 HTML:", html_path, f"({os.path.getsize(html_path)/1024:.0f} KB)")
    if BUG_FILE:
        print("生成 Excel:", xlsx_path)
    else:
        print("未加载 BUG 表，跳过结果 Excel 导出（BUG 数据需在工具中上传）")


if __name__ == "__main__":
    main()
